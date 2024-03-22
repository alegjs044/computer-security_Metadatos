
import os
from PIL import Image
from PIL.ExifTags import TAGS, GPSTAGS, IFD
from docx import Document
from openpyxl import load_workbook
from pdfminer.pdfparser import PDFParser
from pdfminer.pdfdocument import PDFDocument

def get_image_metadata(image_path):
    image_metadata = {}
    image_object = Image.open(image_path)
    exif_data = image_object.getexif()
    if exif_data:
        for tag_id, value in exif_data.items():
            tag = TAGS.get(tag_id)
            image_metadata[tag] = value
            if isinstance(value, bytes):
                value = value.decode()
            print(f"{tag}: {value}")
    return image_metadata

def get_geo(exif):
    gps_info = exif.get(IFD.GPSInfo)
    gps_data = {}
    if isinstance(gps_info, dict): 
        for key, value in gps_info.items():
            tag = GPSTAGS.get(key)
            gps_data[tag] = value
    return gps_data

def get_coordinates(gps_data):
    lat_ref = 1 if gps_data['GPSLatitudeRef'] == 'N' else -1
    long_ref = 1 if gps_data['GPSLongitudeRef'] == 'E' else -1
    latitude = gps_data['GPSLatitude']
    latitude = float(latitude[0] + latitude[1] / 60 + latitude[2] / 3600)
    longitude = gps_data['GPSLongitude']
    longitude = float(longitude[0] + longitude[1] / 60 + longitude[2] / 3600)
    return latitude * lat_ref, longitude * long_ref

def extract_docx_metadata(docx_path):
    docx_metadata = {}
    doc = Document(docx_path)
    core_properties = doc.core_properties
    
    docx_metadata['Title'] = core_properties.title
    docx_metadata['Author'] = core_properties.author
    docx_metadata['Subject'] = core_properties.subject
    docx_metadata['Keywords'] = core_properties.keywords
    docx_metadata['Comments'] = core_properties.comments
    
    for prop, value in docx_metadata.items():
        print(f"{prop}: {value}")
        
    return docx_metadata


def extract_xlsx_metadata(xlsx_path):
    xlsx_metadata = {}
    wb = load_workbook(filename=xlsx_path)
    for prop in wb.properties:
        xlsx_metadata[prop] = str(getattr(wb.properties, prop))
        print(f"{prop}: {xlsx_metadata[prop]}")
    return xlsx_metadata

def extract_pdf_metadata(pdf_path):
    pdf_metadata = {}
    with open(pdf_path, 'rb') as f:
        parser = PDFParser(f)
        doc = PDFDocument(parser)
        info_dict = doc.info[0]  
        for key, value in info_dict.items():
            key_str = key.decode('latin-1') if isinstance(key, bytes) else key
            value_str = value.decode('latin-1') if isinstance(value, bytes) else value
            pdf_metadata[key_str] = value_str
            print(f"{key_str}: {value_str}")
    return pdf_metadata

if __name__ == "__main__":
    directory_path = input("Ingrese la ruta del directorio: ")
    for filename in os.listdir(directory_path):
        file_path = os.path.join(directory_path, filename)
        if filename.lower().endswith(('.jpg', '.jpeg', '.png')):
            print(f"\n∞∞∞ Metadatos de la Imagen: {filename} ∞∞∞")
            get_image_metadata(file_path)
            gps_results = get_geo(Image.open(file_path).getexif())
            if gps_results:
                print("\n∞∞∞ Datos GPS ∞∞∞")
                print(gps_results)
                print("\n∞∞∞ Coordenadas (Latitud-Longitud) de la imagen ∞∞∞")
                print(get_coordinates(gps_results))
        elif filename.lower().endswith('.docx'):
            print(f"\n∞∞∞ Metadatos del Documento Word: {filename} ∞∞∞")
            extract_docx_metadata(file_path)
        elif filename.lower().endswith('.xlsx'):
            print(f"\n∞∞∞ Metadatos de la Hoja de Cálculo Excel: {filename} ∞∞∞")
            extract_xlsx_metadata(file_path)
        elif filename.lower().endswith('.pdf'):
            print(f"\n∞∞∞ Metadatos del Documento PDF: {filename} ∞∞∞")
            extract_pdf_metadata(file_path)
        else:
            print(f"No se pudo extraer metadatos de {filename}: formato no compatible.")

